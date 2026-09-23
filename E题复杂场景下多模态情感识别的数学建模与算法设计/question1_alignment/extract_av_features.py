from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import pickle
import time
from pathlib import Path

import av
import cv2
import numpy as np
from openpyxl import load_workbook


AUDIO_DIM = 74
VISION_DIM = 35
CLASS_MAP = {"Negative": 0, "Neutral": 1, "Positive": 2}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_labels(label_path: Path) -> list[dict]:
    workbook = load_workbook(label_path, read_only=True, data_only=True)
    sheet = workbook["label"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(value) for value in next(rows)]
    records = []
    for row in rows:
        if row[0] is None:
            continue
        record = dict(zip(header, row))
        record["video_id"] = str(record["video_id"])
        record["clip_id"] = str(record["clip_id"])
        record["text"] = str(record["text"] or "")
        record["label"] = float(record["label"])
        record["annotation"] = str(record["annotation"])
        records.append(record)
    workbook.close()
    return records


def media_duration(path: Path) -> float:
    with av.open(str(path)) as container:
        if container.duration:
            return float(container.duration / av.time_base)
        durations = []
        for stream in container.streams:
            if stream.duration is not None and stream.time_base is not None:
                durations.append(float(stream.duration * stream.time_base))
        if not durations:
            raise RuntimeError(f"No duration metadata in {path}")
        return max(durations)


def decode_audio(path: Path, target_rate: int, duration: float) -> tuple[np.ndarray, float]:
    with av.open(str(path)) as container:
        if not container.streams.audio:
            return np.zeros(0, dtype=np.float32), 0.0
        resampler = av.AudioResampler(format="s16", layout="mono", rate=target_rate)
        chunks = []
        for frame in container.decode(audio=0):
            for converted in resampler.resample(frame):
                chunks.append(converted.to_ndarray().reshape(-1))
        for converted in resampler.resample(None):
            chunks.append(converted.to_ndarray().reshape(-1))
    if not chunks:
        return np.zeros(0, dtype=np.float32), 0.0
    audio = np.concatenate(chunks).astype(np.float32) / 32768.0
    expected = max(1, int(round(duration * target_rate)))
    valid_duration = min(len(audio), expected) / target_rate
    if len(audio) < expected:
        audio = np.pad(audio, (0, expected - len(audio)))
    else:
        audio = audio[:expected]
    return audio, valid_duration


def hz_to_mel(hz: np.ndarray) -> np.ndarray:
    return 2595.0 * np.log10(1.0 + hz / 700.0)


def mel_to_hz(mel: np.ndarray) -> np.ndarray:
    return 700.0 * (10.0 ** (mel / 2595.0) - 1.0)


def mel_filterbank(sample_rate: int, n_fft: int, n_mels: int = 26) -> np.ndarray:
    mel_points = np.linspace(hz_to_mel(np.array([0.0]))[0], hz_to_mel(np.array([sample_rate / 2]))[0], n_mels + 2)
    hz_points = mel_to_hz(mel_points)
    bins = np.floor((n_fft + 1) * hz_points / sample_rate).astype(int)
    bins = np.clip(bins, 0, n_fft // 2)
    filters = np.zeros((n_mels, n_fft // 2 + 1), dtype=np.float64)
    for index in range(n_mels):
        left, center, right = bins[index : index + 3]
        if center <= left:
            center = min(left + 1, n_fft // 2)
        if right <= center:
            right = min(center + 1, n_fft // 2)
        if center > left:
            filters[index, left:center] = (np.arange(left, center) - left) / (center - left)
        if right > center:
            filters[index, center:right] = (right - np.arange(center, right)) / (right - center)
    return filters


def frame_signal(signal: np.ndarray, frame_length: int, hop_length: int) -> np.ndarray:
    if len(signal) < frame_length:
        signal = np.pad(signal, (0, frame_length - len(signal)))
    remainder = (len(signal) - frame_length) % hop_length
    if remainder:
        signal = np.pad(signal, (0, hop_length - remainder))
    windows = np.lib.stride_tricks.sliding_window_view(signal, frame_length)[::hop_length]
    return np.asarray(windows, dtype=np.float64)


def audio_bin_features(signal: np.ndarray, sample_rate: int) -> np.ndarray:
    frame_length = max(64, int(round(0.025 * sample_rate)))
    hop_length = max(32, int(round(0.010 * sample_rate)))
    n_fft = 512
    frames = frame_signal(signal.astype(np.float64), frame_length, hop_length)
    frames = frames - frames.mean(axis=1, keepdims=True)
    windowed = frames * np.hanning(frame_length)[None, :]
    spectrum = np.abs(np.fft.rfft(windowed, n=n_fft, axis=1))
    power = spectrum**2 + 1e-12
    frequencies = np.fft.rfftfreq(n_fft, 1.0 / sample_rate)

    filters = mel_filterbank(sample_rate, n_fft, 26)
    log_mel = np.log(np.maximum(power @ filters.T, 1e-12))
    indices = np.arange(20)[:, None]
    mel_indices = np.arange(26)[None, :]
    dct = np.cos(np.pi / 26.0 * (mel_indices + 0.5) * indices)
    mfcc = log_mel @ dct.T
    mfcc_stats = np.concatenate([mfcc.mean(axis=0), mfcc.std(axis=0)])

    chroma = np.zeros((power.shape[0], 12), dtype=np.float64)
    valid_frequency = frequencies > 27.5
    midi = np.zeros_like(frequencies, dtype=int)
    midi[valid_frequency] = np.rint(69 + 12 * np.log2(frequencies[valid_frequency] / 440.0)).astype(int)
    pitch_class = np.mod(midi, 12)
    for pitch in range(12):
        mask = valid_frequency & (pitch_class == pitch)
        if np.any(mask):
            chroma[:, pitch] = power[:, mask].sum(axis=1)
    chroma /= np.maximum(chroma.sum(axis=1, keepdims=True), 1e-12)
    chroma_mean = chroma.mean(axis=0)

    total_power = power.sum(axis=1) + 1e-12
    centroid = (power * frequencies[None, :]).sum(axis=1) / total_power
    bandwidth = np.sqrt((power * (frequencies[None, :] - centroid[:, None]) ** 2).sum(axis=1) / total_power)
    cumulative = np.cumsum(power, axis=1)
    rolloff_index = np.argmax(cumulative >= 0.85 * total_power[:, None], axis=1)
    rolloff = frequencies[rolloff_index]
    flatness = np.exp(np.mean(np.log(power), axis=1)) / np.mean(power, axis=1)
    probability = power / total_power[:, None]
    entropy = -np.sum(probability * np.log(probability + 1e-12), axis=1) / np.log(power.shape[1])
    normalized_spectrum = spectrum / np.maximum(spectrum.sum(axis=1, keepdims=True), 1e-12)
    flux = np.sqrt(np.mean(np.diff(normalized_spectrum, axis=0, prepend=normalized_spectrum[:1]) ** 2, axis=1))
    low_ratio = power[:, frequencies < 300].sum(axis=1) / total_power
    high_ratio = power[:, frequencies > 3000].sum(axis=1) / total_power
    spectral_matrix = np.column_stack(
        [
            centroid / (sample_rate / 2),
            bandwidth / (sample_rate / 2),
            rolloff / (sample_rate / 2),
            flatness,
            entropy,
            flux,
            low_ratio,
            high_ratio,
        ]
    )
    spectral_stats = np.concatenate([spectral_matrix.mean(axis=0), spectral_matrix.std(axis=0)])

    rms = np.sqrt(np.mean(frames**2, axis=1) + 1e-12)
    zcr = np.mean(np.abs(np.diff(np.signbit(frames), axis=1)), axis=1)
    pitch_band = (frequencies >= 70) & (frequencies <= 400)
    if np.any(pitch_band):
        local = power[:, pitch_band]
        pitch_frequencies = frequencies[pitch_band][np.argmax(local, axis=1)]
        pitch_frequencies = np.where(rms > 1e-4, pitch_frequencies, 0.0)
    else:
        pitch_frequencies = np.zeros(len(frames))
    voiced = pitch_frequencies[pitch_frequencies > 0]
    pitch_stats = np.array([voiced.mean(), voiced.std()]) if len(voiced) else np.zeros(2)
    temporal_stats = np.array([zcr.mean(), zcr.std(), rms.mean(), rms.std()])

    result = np.concatenate([mfcc_stats, chroma_mean, spectral_stats, pitch_stats, temporal_stats])
    if result.shape != (AUDIO_DIM,):
        raise AssertionError(f"Audio feature dimension is {result.shape}, expected {(AUDIO_DIM,)}")
    return np.nan_to_num(result, nan=0.0, posinf=0.0, neginf=0.0).astype(np.float32)


def extract_audio(path: Path, duration: float, num_bins: int, sample_rate: int) -> tuple[np.ndarray, np.ndarray, float]:
    audio, valid_duration = decode_audio(path, sample_rate, duration)
    features = np.zeros((num_bins, AUDIO_DIM), dtype=np.float32)
    mask = np.zeros(num_bins, dtype=np.uint8)
    if len(audio) == 0:
        return features, mask, valid_duration
    for index in range(num_bins):
        start = int(round(index * duration / num_bins * sample_rate))
        end = int(round((index + 1) * duration / num_bins * sample_rate))
        if start >= int(valid_duration * sample_rate) or end <= start:
            continue
        segment = audio[start:end]
        if len(segment) < 32:
            continue
        features[index] = audio_bin_features(segment, sample_rate)
        mask[index] = 1
    return features, mask, valid_duration


def resize_frame(frame: np.ndarray, max_width: int) -> np.ndarray:
    height, width = frame.shape[:2]
    if width <= max_width:
        return frame
    scale = max_width / width
    return cv2.resize(frame, (max_width, max(1, int(round(height * scale)))), interpolation=cv2.INTER_AREA)


def largest_face(detector: cv2.CascadeClassifier, gray: np.ndarray) -> tuple[int, int, int, int] | None:
    faces = detector.detectMultiScale(gray, scaleFactor=1.1, minNeighbors=5, minSize=(24, 24))
    if len(faces) == 0:
        return None
    return tuple(int(value) for value in max(faces, key=lambda item: item[2] * item[3]))


def static_vision_features(
    bgr: np.ndarray, detector: cv2.CascadeClassifier
) -> tuple[np.ndarray, np.ndarray, tuple[int, int, int, int] | None]:
    gray_u8 = cv2.cvtColor(bgr, cv2.COLOR_BGR2GRAY)
    gray = gray_u8.astype(np.float32) / 255.0
    height, width = gray.shape
    face = largest_face(detector, gray_u8)
    if face is None:
        face_values = [0.0] * 6
        x, y = int(0.2 * width), int(0.15 * height)
        fw, fh = max(1, int(0.6 * width)), max(1, int(0.7 * height))
    else:
        x, y, fw, fh = face
        face_values = [1.0, (x + fw / 2) / width, (y + fh / 2) / height, fw / width, fh / height, fw * fh / (width * height)]

    intensity = [
        float(gray.mean()),
        float(gray.std()),
        *[float(value) for value in np.percentile(gray, [10, 50, 90])],
        float(np.mean(gray < 0.2)),
    ]
    hsv = cv2.cvtColor(bgr, cv2.COLOR_BGR2HSV).astype(np.float32)
    hsv[:, :, 0] /= 179.0
    hsv[:, :, 1:] /= 255.0
    color = [float(hsv[:, :, channel].mean()) for channel in range(3)] + [float(hsv[:, :, channel].std()) for channel in range(3)]

    edges = cv2.Canny(gray_u8, 80, 160)
    gx = cv2.Sobel(gray, cv2.CV_32F, 1, 0, ksize=3)
    gy = cv2.Sobel(gray, cv2.CV_32F, 0, 1, ksize=3)
    magnitude = np.sqrt(gx**2 + gy**2)
    laplacian = cv2.Laplacian(gray, cv2.CV_32F)
    histogram = np.histogram(gray, bins=32, range=(0.0, 1.0))[0].astype(np.float64)
    histogram /= max(histogram.sum(), 1.0)
    gray_entropy = -np.sum(histogram * np.log(histogram + 1e-12)) / np.log(32)
    texture = [float(np.mean(edges > 0)), float(magnitude.mean()), float(magnitude.std()), float(laplacian.var()), float(gray_entropy)]

    x2, y2 = min(width, x + fw), min(height, y + fh)
    region = gray[max(0, y) : y2, max(0, x) : x2]
    if region.size == 0:
        region = gray
    thirds = np.array_split(region, 3, axis=0)
    regional = []
    for third in thirds:
        regional.extend([float(third.mean()), float(third.std())])
    midpoint = region.shape[1] // 2
    left = region[:, :midpoint]
    right = region[:, region.shape[1] - midpoint :]
    if midpoint and left.size and right.size:
        regional.append(float(abs(left.mean() - right.mean())))
        regional.append(float(np.mean(np.abs(left - np.fliplr(right)))))
    else:
        regional.extend([0.0, 0.0])
    static = np.asarray(face_values + intensity + color + texture + regional, dtype=np.float32)
    if static.shape != (31,):
        raise AssertionError(f"Static vision feature dimension is {static.shape}, expected (31,)")
    return static, gray_u8, face


def extract_vision(path: Path, duration: float, num_bins: int, max_width: int) -> tuple[np.ndarray, np.ndarray, np.ndarray, int]:
    centers = (np.arange(num_bins, dtype=np.float64) + 0.5) * duration / num_bins
    best_distance = np.full(num_bins, np.inf)
    selected_frames: list[np.ndarray | None] = [None] * num_bins
    selected_times = np.full(num_bins, np.nan, dtype=np.float64)
    decoded_frames = 0
    with av.open(str(path)) as container:
        stream = container.streams.video[0]
        fps = float(stream.average_rate) if stream.average_rate else 25.0
        for frame_index, frame in enumerate(container.decode(video=0)):
            decoded_frames += 1
            timestamp = float(frame.time) if frame.time is not None else frame_index / fps
            bin_index = int(np.clip(np.rint(timestamp / duration * num_bins - 0.5), 0, num_bins - 1))
            distance = abs(timestamp - centers[bin_index])
            if distance < best_distance[bin_index]:
                bgr = resize_frame(frame.to_ndarray(format="bgr24"), max_width)
                selected_frames[bin_index] = bgr
                selected_times[bin_index] = timestamp
                best_distance[bin_index] = distance

    detector_path = cv2.data.haarcascades + "haarcascade_frontalface_default.xml"
    detector = cv2.CascadeClassifier(detector_path)
    features = np.zeros((num_bins, VISION_DIM), dtype=np.float32)
    mask = np.zeros(num_bins, dtype=np.uint8)
    previous_gray = None
    for index, frame in enumerate(selected_frames):
        if frame is None:
            previous_gray = None
            continue
        static, gray, _ = static_vision_features(frame, detector)
        motion = np.zeros(4, dtype=np.float32)
        if previous_gray is not None and previous_gray.shape == gray.shape:
            flow = cv2.calcOpticalFlowFarneback(previous_gray, gray, None, 0.5, 3, 15, 3, 5, 1.2, 0)
            dx = flow[:, :, 0]
            dy = flow[:, :, 1]
            magnitude = np.sqrt(dx**2 + dy**2)
            scale = max(gray.shape)
            motion = np.array([dx.mean() / scale, dy.mean() / scale, magnitude.mean() / scale, magnitude.std() / scale], dtype=np.float32)
        features[index] = np.concatenate([static[:23], motion, static[23:]])
        mask[index] = 1
        previous_gray = gray
    if features.shape[1] != VISION_DIM:
        raise AssertionError(f"Vision feature dimension is {features.shape}")
    return features, mask, selected_times, decoded_frames


def build_time_bounds(duration: float, num_bins: int) -> np.ndarray:
    edges = np.linspace(0.0, duration, num_bins + 1, dtype=np.float64)
    return np.column_stack([edges[:-1], edges[1:]]).astype(np.float32)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract aligned audio and visual features for question 1.")
    parser.add_argument("--data-root", type=Path, required=True, help="Path to E题数据/E题数据")
    parser.add_argument("--work-dir", type=Path, required=True)
    parser.add_argument("--num-bins", type=int, default=50)
    parser.add_argument("--sample-rate", type=int, default=16000)
    parser.add_argument("--vision-max-width", type=int, default=320)
    args = parser.parse_args()

    raw_root = args.data_root / "附件1-数据集原始多模态样本" / "MOSEI数据集部分原始视频-100条"
    label_path = raw_root / "label-100.xlsx"
    records = load_labels(label_path)
    args.work_dir.mkdir(parents=True, exist_ok=True)
    count = len(records)
    if count != 100:
        raise RuntimeError(f"Expected 100 labels, found {count}")

    audio = np.zeros((count, args.num_bins, AUDIO_DIM), dtype=np.float32)
    vision = np.zeros((count, args.num_bins, VISION_DIM), dtype=np.float32)
    audio_mask = np.zeros((count, args.num_bins), dtype=np.uint8)
    vision_mask = np.zeros((count, args.num_bins), dtype=np.uint8)
    time_bounds = np.zeros((count, args.num_bins, 2), dtype=np.float32)
    vision_times = np.full((count, args.num_bins), np.nan, dtype=np.float32)
    durations = np.zeros(count, dtype=np.float32)
    audio_durations = np.zeros(count, dtype=np.float32)
    ids = []
    source_files = []
    log_records = []

    manifest_path = args.work_dir / "feature_manifest.csv"
    manifest_fields = [
        "sample_id",
        "video_id",
        "clip_id",
        "modality",
        "source_file",
        "duration_s",
        "valid_duration_s",
        "feature_shape",
        "feature_dim",
        "num_bins",
        "bin_width_s",
        "valid_bins",
        "padding_rule",
        "extractor",
        "source_sha256",
    ]
    with manifest_path.open("w", encoding="utf-8-sig", newline="") as manifest_handle, (args.work_dir / "processing_log.jsonl").open(
        "w", encoding="utf-8"
    ) as log_handle:
        manifest = csv.DictWriter(manifest_handle, fieldnames=manifest_fields)
        manifest.writeheader()
        for index, record in enumerate(records):
            started = time.perf_counter()
            sample_id = f"{record['video_id']}$_${record['clip_id']}"
            video_path = raw_root / record["video_id"] / f"{record['clip_id']}.mp4"
            if not video_path.exists():
                raise FileNotFoundError(video_path)
            duration = media_duration(video_path)
            source_hash = sha256_file(video_path)
            sample_audio, sample_audio_mask, audio_duration = extract_audio(video_path, duration, args.num_bins, args.sample_rate)
            sample_vision, sample_vision_mask, sample_vision_times, decoded_frames = extract_vision(
                video_path, duration, args.num_bins, args.vision_max_width
            )
            audio[index] = sample_audio
            vision[index] = sample_vision
            audio_mask[index] = sample_audio_mask
            vision_mask[index] = sample_vision_mask
            time_bounds[index] = build_time_bounds(duration, args.num_bins)
            vision_times[index] = sample_vision_times
            durations[index] = duration
            audio_durations[index] = audio_duration
            ids.append(sample_id)
            source_files.append(str(video_path.relative_to(args.data_root)))

            elapsed = time.perf_counter() - started
            log_record = {
                "sample_index": index,
                "sample_id": sample_id,
                "source_file": source_files[-1],
                "duration_s": duration,
                "audio_valid_duration_s": audio_duration,
                "audio_valid_bins": int(sample_audio_mask.sum()),
                "vision_valid_bins": int(sample_vision_mask.sum()),
                "decoded_video_frames": decoded_frames,
                "elapsed_s": elapsed,
                "status": "ok",
            }
            log_records.append(log_record)
            log_handle.write(json.dumps(log_record, ensure_ascii=False) + "\n")
            print(f"[{index + 1:03d}/{count:03d}] {sample_id}  T={duration:.3f}s  A={sample_audio_mask.sum():02d}  V={sample_vision_mask.sum():02d}")

            common = {
                "sample_id": sample_id,
                "video_id": record["video_id"],
                "clip_id": record["clip_id"],
                "source_file": source_files[-1],
                "duration_s": f"{duration:.6f}",
                "num_bins": args.num_bins,
                "bin_width_s": f"{duration / args.num_bins:.6f}",
                "source_sha256": source_hash,
            }
            manifest.writerow(
                {
                    **common,
                    "modality": "text",
                    "valid_duration_s": f"{duration:.6f}",
                    "feature_shape": f"{args.num_bins}x768",
                    "feature_dim": 768,
                    "valid_bins": args.num_bins if record["text"].strip() else 0,
                    "padding_rule": "zero vector with text_mask=0",
                    "extractor": "distilbert-base-uncased; mean of last 4 hidden layers; overlap-weighted pooling",
                }
            )
            manifest.writerow(
                {
                    **common,
                    "modality": "audio",
                    "valid_duration_s": f"{audio_duration:.6f}",
                    "feature_shape": f"{args.num_bins}x{AUDIO_DIM}",
                    "feature_dim": AUDIO_DIM,
                    "valid_bins": int(sample_audio_mask.sum()),
                    "padding_rule": "zero vector with audio_mask=0",
                    "extractor": "16 kHz mono; MFCC/chroma/spectral/pitch/energy statistics",
                }
            )
            manifest.writerow(
                {
                    **common,
                    "modality": "vision",
                    "valid_duration_s": f"{duration:.6f}",
                    "feature_shape": f"{args.num_bins}x{VISION_DIM}",
                    "feature_dim": VISION_DIM,
                    "valid_bins": int(sample_vision_mask.sum()),
                    "padding_rule": "nearest center frame; zero vector with vision_mask=0",
                    "extractor": "OpenCV face/appearance/texture/optical-flow descriptors",
                }
            )

    payload = {
        "id": ids,
        "video_id": [record["video_id"] for record in records],
        "clip_id": [record["clip_id"] for record in records],
        "raw_text": [record["text"] for record in records],
        "annotations": [record["annotation"] for record in records],
        "classification_labels": np.asarray([CLASS_MAP[record["annotation"]] for record in records], dtype=np.int64),
        "regression_labels": np.asarray([record["label"] for record in records], dtype=np.float32),
        "audio": audio,
        "vision": vision,
        "audio_mask": audio_mask,
        "vision_mask": vision_mask,
        "time_bounds": time_bounds,
        "vision_timestamps": vision_times,
        "durations": durations,
        "audio_valid_durations": audio_durations,
        "source_file": source_files,
        "settings": {
            "num_bins": args.num_bins,
            "audio_sample_rate": args.sample_rate,
            "vision_max_width": args.vision_max_width,
        },
    }
    with (args.work_dir / "base_features.pkl").open("wb") as handle:
        pickle.dump(payload, handle, protocol=4)
    samples = [
        {"sample_index": index, "sample_id": ids[index], "text": records[index]["text"], "duration_s": float(durations[index])}
        for index in range(count)
    ]
    (args.work_dir / "samples.json").write_text(json.dumps(samples, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Saved base features for {count} samples to {args.work_dir}")


if __name__ == "__main__":
    main()
